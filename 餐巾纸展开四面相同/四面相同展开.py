#!/usr/bin/env python3
"""
餐巾纸展开图片处理脚本
====================

此脚本自动完成以下任务：
1. 从"原图"文件夹读取所有图片
2. 将单张图片展开为完整的4倍大小餐巾纸（上面两个角旋转180度）
3. 使用展开后的图片替换PowerPoint模板中的图片
4. 将结果保存到"成果"文件夹中

展开逻辑：
- 左上角：原图旋转180度
- 右上角：原图旋转180度  
- 左下角：原图复制
- 右下角：原图复制
（上面两个象限旋转180度，下面两个象限保持原样）

专为非技术用户设计 - 只需点击运行！
"""

import os
import sys
import shutil
from PIL import Image
from pptx import Presentation
from pathlib import Path

# 支持的图片格式
SUPPORTED_EXTENSIONS = [".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".webp"]

# 文件夹名称
INPUT_FOLDER = "原图"
OUTPUT_FOLDER = "成果"
TEMPLATE_FILE = "单页模板.pptx"

# 图片到PowerPoint形状名称的映射
# 每个PPTX文件将使用两张图片：原图和展开图
IMAGE_SHAPE_MAPPING = {
    "original": ["1","3"],      # 原图替换的形状名称列表
    "unfolded": ["2"]       # 展开图替换的形状名称列表
}

# TODO: 根据您的PowerPoint模板调整上述映射
# 示例：
# IMAGE_SHAPE_MAPPING = {
#     "original": ["原图形状", "另一个原图形状"],
#     "unfolded": ["展开图形状", "另一个展开图形状"] 
# }

def print_separator():
    """打印视觉分隔线以提高可读性"""
    print("=" * 60)

def get_script_directory():
    """获取脚本所在目录"""
    if getattr(sys, 'frozen', False):
        # 如果作为exe运行（使用PyInstaller编译）
        return Path(sys.executable).parent
    else:
        # 如果作为Python脚本运行
        return Path(__file__).parent

def find_images_in_folder(folder_path):
    """查找文件夹中的所有图片文件"""
    image_files = []
    
    if not folder_path.exists():
        return image_files
    
    for ext in SUPPORTED_EXTENSIONS:
        image_files.extend(list(folder_path.glob(f"*{ext}")))
        image_files.extend(list(folder_path.glob(f"*{ext.upper()}")))
    
    return image_files

def create_four_same_napkin(image_path, output_path):
    """
    从单张图片创建完整的展开餐巾纸（上面两个角旋转180度）
    逻辑：
    - 左上角：原图旋转180度
    - 右上角：原图旋转180度
    - 左下角：原图复制
    - 右下角：原图复制
    
    Args:
        image_path: 输入图片路径
        output_path: 输出图片路径
    
    Returns:
        bool: 处理是否成功
    """
    try:
        # 打开原始图片
        with Image.open(image_path) as original:
            print(f"  📐 原图尺寸：{original.size[0]}x{original.size[1]}像素")
            
            # 获取原图尺寸
            orig_width, orig_height = original.size
            
            # 创建4倍大小的新画布
            new_width = orig_width * 2
            new_height = orig_height * 2
            new_image = Image.new('RGB', (new_width, new_height), 'white')
            
            print(f"  🔧 创建展开画布：{new_width}x{new_height}像素")
            
            # 创建旋转180度的版本
            rotated_image = original.rotate(180)
            
            # 四个象限：上面两个旋转180度，下面两个保持原样
            # 左上角 - 旋转180度
            new_image.paste(rotated_image, (0, 0))
            print(f"  📍 左上角：原图旋转180度")
            
            # 右上角 - 旋转180度
            new_image.paste(rotated_image, (orig_width, 0))
            print(f"  📍 右上角：原图旋转180度")
            
            # 左下角 - 原图复制
            new_image.paste(original, (0, orig_height))
            print(f"  📍 左下角：原图复制")
            
            # 右下角 - 原图复制
            new_image.paste(original, (orig_width, orig_height))
            print(f"  📍 右下角：原图复制")
            
            # 保存展开的餐巾纸图片
            new_image.save(output_path, quality=95, optimize=True)
            print(f"  💾 已保存展开图：{output_path.name}")
            
            return True
            
    except Exception as e:
        print(f"  ❌ 处理图片时出错：{e}")
        return False

def replace_images_in_presentation(pptx_path, original_image_path, unfolded_image_path, image_name=None):
    """
    替换PowerPoint演示文稿中的图片
    使用硬编码映射替换原图和展开图到指定的形状
    
    Args:
        pptx_path: PowerPoint文件路径
        original_image_path: 原始图片路径
        unfolded_image_path: 展开图片路径
        image_name: 图片名称（用于文本替换，可选）
    
    Returns:
        bool: 是否成功替换
    """
    try:
        prs = Presentation(pptx_path)
        replacements_made = 0
        text_replacements_made = 0
        
        print(f"  🔍 正在根据映射替换图片形状...")
        print(f"  📷 原图: {original_image_path.name}")
        print(f"  📷 展开图: {unfolded_image_path.name}")
        
        # 创建图片类型到文件路径的映射
        image_files = {
            "original": original_image_path,
            "unfolded": unfolded_image_path
        }
        
        # 处理图片替换
        for slide_idx, slide in enumerate(prs.slides):
            print(f"  📄 正在处理幻灯片 {slide_idx + 1}")
            
            # 遍历每种图片类型的映射
            for image_type, shape_names in IMAGE_SHAPE_MAPPING.items():
                image_path = image_files.get(image_type)
                
                if not image_path or not image_path.exists():
                    print(f"    ⚠️  跳过{image_type}图片（文件不存在）")
                    continue
                
                # 对于每个形状名称列表中的形状
                for shape_name in shape_names:
                    print(f"    🔍 查找形状: '{shape_name}' (用于{image_type}图片)")
                    
                    # 按名称查找形状
                    shape_found = False
                    for shape in slide.shapes:
                        if hasattr(shape, 'name') and shape.name == shape_name and shape.shape_type == 13:  # 13表示图片类型
                            print(f"    🔄 正在替换形状'{shape_name}'为{image_path.name}")
                            
                            # 保存原始位置、大小和索引
                            left = shape.left
                            top = shape.top
                            width = shape.width
                            height = shape.height
                            shape_index = slide.shapes.index(shape)
                            
                            # 移除原始形状
                            sp = shape._element
                            sp.getparent().remove(sp)
                            
                            # 添加新图片并保持原始位置和大小
                            new_pic = slide.shapes.add_picture(str(image_path), left, top, width, height)
                            
                            # 确保新图片具有与原始图片相同的属性
                            new_pic.left = left
                            new_pic.top = top
                            new_pic.width = width
                            new_pic.height = height
                            
                            # 正确处理XML元素顺序
                            new_pic_element = new_pic._element
                            slide.shapes._spTree.remove(new_pic_element)
                            
                            # 确保插入索引不会破坏spTree结构
                            spTree = slide.shapes._spTree
                            group_props_count = 0
                            
                            # 计算组属性元素的数量
                            for child in spTree:
                                if child.tag.endswith('}nvGrpSpPr') or child.tag.endswith('}grpSpPr'):
                                    group_props_count += 1
                                else:
                                    break
                            
                            correct_insert_index = max(group_props_count, shape_index + group_props_count)
                            
                            if correct_insert_index > len(spTree):
                                spTree.append(new_pic_element)
                            else:
                                spTree.insert(correct_insert_index, new_pic_element)
                            
                            replacements_made += 1
                            print(f"    ✅ 成功替换形状'{shape_name}'")
                            shape_found = True
                            break
                    
                    if not shape_found:
                        print(f"    ❌ 未找到形状: '{shape_name}'")
        
        # 处理文本替换（如果提供了图片名称）
        if image_name:
            print(f"  🔍 正在查找文本框进行替换...")
            
            for slide_idx, slide in enumerate(prs.slides):
                # 查找文本框并替换内容
                for shape in slide.shapes:
                    if hasattr(shape, 'text_frame') and shape.text_frame:
                        # 检查是否有文本内容
                        if shape.text_frame.text.strip():
                            print(f"    📝 正在替换文本为'{image_name}'")
                            
                            # 保存原始格式并替换文本
                            text_frame = shape.text_frame
                            
                            # 如果文本框有内容，保持第一段的格式
                            if text_frame.paragraphs:
                                # 获取第一段
                                first_paragraph = text_frame.paragraphs[0]
                                
                                # 保存原始段落格式
                                original_alignment = first_paragraph.alignment if hasattr(first_paragraph, 'alignment') else None
                                
                                # 如果第一段有运行（runs），保存第一个运行的格式
                                original_font = None
                                if first_paragraph.runs:
                                    first_run = first_paragraph.runs[0]
                                    original_font = {
                                        'name': first_run.font.name,
                                        'size': first_run.font.size,
                                        'bold': first_run.font.bold,
                                        'italic': first_run.font.italic,
                                        'color': first_run.font.color.rgb if hasattr(first_run.font.color, 'rgb') else None
                                    }
                                
                                # 清除现有内容
                                text_frame.clear()
                                
                                # 添加新段落
                                new_paragraph = text_frame.paragraphs[0]
                                
                                # 恢复段落格式
                                if original_alignment is not None:
                                    new_paragraph.alignment = original_alignment
                                
                                # 添加新文本运行
                                new_run = new_paragraph.add_run()
                                new_run.text = image_name
                                
                                # 恢复字体格式
                                if original_font:
                                    if original_font['name']:
                                        new_run.font.name = original_font['name']
                                    if original_font['size']:
                                        new_run.font.size = original_font['size']
                                    if original_font['bold'] is not None:
                                        new_run.font.bold = original_font['bold']
                                    if original_font['italic'] is not None:
                                        new_run.font.italic = original_font['italic']
                                    if original_font['color']:
                                        new_run.font.color.rgb = original_font['color']
                                
                                text_replacements_made += 1
                                print(f"    ✅ 成功更新文本框（已保持格式）")
                                break  # 只替换第一个找到的文本框
        
        # 保存演示文稿
        prs.save(pptx_path)
        print(f"  📊 替换统计:")
        print(f"     - 图片替换: {replacements_made} 个")
        print(f"     - 文本替换: {text_replacements_made} 个")
        
        return replacements_made > 0
        
    except Exception as e:
        print(f"  ❌ 更新演示文稿时出错：{e}")
        return False

def process_image(image_path, template_path, output_folder):
    """处理单张图片：创建展开版本并生成PPTX文件"""
    print(f"\n📷 正在处理图片：{image_path.name}")
    
    # 生成PPTX文件名并检查是否已存在
    base_name = image_path.stem
    pptx_name = f"{base_name}.pptx"
    pptx_path = output_folder / pptx_name
    
    # 检查输出文件是否已存在
    if pptx_path.exists():
        print(f"  ⏭️  跳过处理（输出文件已存在）: {pptx_name}")
        return "skipped"  # 返回"skipped"表示跳过
    
    # 生成临时展开图片文件名
    temp_image_name = f"{base_name}_unfolded.png"
    temp_image_path = output_folder / temp_image_name
    
    # 创建展开的餐巾纸图片
    print(f"  🔧 正在展开餐巾纸（上面两角旋转180度）...")
    success = create_four_same_napkin(image_path, temp_image_path)
    
    if not success:
        print(f"  ❌ 展开失败：{image_path.name}")
        return False
    
    # 复制模板到输出目录
    try:
        shutil.copy2(template_path, pptx_path)
        print(f"  📋 已创建PPTX副本：{pptx_name}")
    except Exception as e:
        print(f"  ❌ 复制模板失败：{e}")
        # 删除临时图片文件
        if temp_image_path.exists():
            temp_image_path.unlink()
        return False
    
    # 替换PPTX中的图片（使用原图和展开图）
    print(f"  🖼️  正在更新PPTX中的图片...")
    pptx_success = replace_images_in_presentation(pptx_path, image_path, temp_image_path, base_name)
    
    if pptx_success:
        print(f"  🎉 成功处理图片：{image_path.name}")
        # 删除临时展开图片文件（保留PPTX）
        if temp_image_path.exists():
            temp_image_path.unlink()
            print(f"  🗑️  已删除临时图片：{temp_image_name}")
        return True
    else:
        print(f"  ❌ PPTX处理失败：{image_path.name}")
        # 删除失败的文件
        if temp_image_path.exists():
            temp_image_path.unlink()
        if pptx_path.exists():
            pptx_path.unlink()
        return False

def main():
    """主函数，协调整个处理过程"""
    try:
        # 自动检测工作目录
        working_dir = get_script_directory()
        input_folder = working_dir / INPUT_FOLDER
        output_folder = working_dir / OUTPUT_FOLDER
        template_path = working_dir / TEMPLATE_FILE
        
        print_separator()
        print("🧾 餐巾纸展开处理器")
        print_separator()
        print("此脚本将自动处理输入文件夹中的所有图片，")
        print("并将图片展开为完整的4倍大小餐巾纸图案（上面两个角旋转180度），")
        print("然后生成对应的PowerPoint演示文稿。")
        print(f"工作目录：{working_dir}")
        print(f"输入文件夹：{INPUT_FOLDER}")
        print(f"输出文件夹：{OUTPUT_FOLDER}")
        print(f"模板文件：{TEMPLATE_FILE}")
        print()
        print("📋 展开模式说明：")
        print("  上面旋转180度展开逻辑：")
        print("    ┌─────────┬─────────┐")
        print("    │ 左上角  │ 右上角  │")
        print("    │旋转180度│旋转180度│")
        print("    ├─────────┼─────────┤")
        print("    │ 左下角  │ 右下角  │")
        print("    │ 原图复制 │ 原图复制 │")
        print("    └─────────┴─────────┘")
        print()
        print("🔧 图片映射配置：")
        print("  每个PPTX将使用两张图片：")
        for image_type, shape_names in IMAGE_SHAPE_MAPPING.items():
            image_desc = "原始图片" if image_type == "original" else "展开图片"
            print(f"  - {image_desc}: 替换形状 {shape_names}")
        print("  ※ 可在脚本顶部修改 IMAGE_SHAPE_MAPPING 配置")
        print()
        print(f"📸 支持的图片格式：{', '.join(SUPPORTED_EXTENSIONS)}")
        print()
        
        # 检查模板文件是否存在
        if not template_path.exists():
            print(f"❌ 未找到模板文件：{TEMPLATE_FILE}")
            print("请确保模板PowerPoint文件与此脚本在同一文件夹中。")
            input("按回车键退出...")
            return
        
        # 创建输入文件夹（如果不存在）
        if not input_folder.exists():
            input_folder.mkdir(exist_ok=True)
            print(f"✅ 已创建输入文件夹：{INPUT_FOLDER}")
            print("请将要处理的图片放入此文件夹中。")
            input("按回车键退出...")
            return
        
        # 创建输出文件夹（如果不存在）
        output_folder.mkdir(exist_ok=True)
        
        print_separator()
        print("🔍 扫描输入文件夹")
        print_separator()
        
        # 查找输入文件夹中的所有图片
        image_files = find_images_in_folder(input_folder)
        
        if not image_files:
            print(f"❌ 在输入文件夹「{INPUT_FOLDER}」中未找到图片文件。")
            print("请确保：")
            print("  1. 图片文件放在正确的输入文件夹中")
            print(f"  2. 图片格式为：{', '.join(SUPPORTED_EXTENSIONS)}")
            input("按回车键退出...")
            return
        
        print(f"✅ 找到{len(image_files)}张图片")
        for img_file in image_files:
            print(f"  📷 {img_file.name}")
        
        # 检查现有输出文件
        existing_pptx = list(output_folder.glob("*.pptx"))
        if existing_pptx:
            print(f"\n📄 输出文件夹中已有{len(existing_pptx)}个PPTX文件：")
            for pptx_file in existing_pptx:
                print(f"  📋 {pptx_file.name}")
            print("  ※ 已存在的文件将被跳过")
        
        # 确认开始处理
        print()
        response = input("🚀 准备开始处理？（按回车继续，输入'n'取消）：").lower()
        if response == 'n':
            print("操作已取消。")
            return
        
        print()
        print_separator()
        print("🚀 开始处理")
        print_separator()
        
        # 处理每张图片
        successful = 0
        skipped = 0
        failed = 0
        
        for image_path in image_files:
            result = process_image(image_path, template_path, output_folder)
            if result == "skipped":
                skipped += 1
            elif result == True:
                successful += 1
            else:
                failed += 1
        
        # 最终总结
        print()
        print_separator()
        print("✨ 处理完成")
        print_separator()
        print(f"总图片数：{len(image_files)}")
        print(f"成功处理：{successful}张")
        print(f"跳过处理：{skipped}张（已存在）")
        print(f"处理失败：{failed}张")
        
        if successful > 0 or skipped > 0:
            print(f"\n🎉 您的餐巾纸展开图和演示文稿已准备就绪！")
            print(f"处理结果保存在：{OUTPUT_FOLDER}")
            print("每个图片生成：")
            print("  📋 PowerPoint演示文稿：[原文件名].pptx")
            print()
            print("📝 说明：")
            print("   每个PPTX文件包含原始图片和对应的展开版本")
            print("   根据映射配置替换PowerPoint模板中的不同形状")
            print("   上面两个象限旋转180度，下面两个象限保持原样")
            if skipped > 0:
                print(f"\n⏭️  跳过了{skipped}个已存在的文件以避免重复处理")
        else:
            print("\n⚠️  未成功处理任何图片。")
            print("请检查您的图片文件和模板文件是否正确。")
        
        # 保持窗口打开
        print()
        input("按回车键退出...")
        
    except KeyboardInterrupt:
        print("\n\n⏹️  用户取消操作。")
        input("按回车键退出...")
    except Exception as e:
        print(f"\n❌ 意外错误：{e}")
        print("请检查您的文件并重试。")
        input("按回车键退出...")

if __name__ == "__main__":
    main() 